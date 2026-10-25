import asyncio
import io
from datetime import date
from types import SimpleNamespace

import openpyxl
import pytest
from fastapi import HTTPException
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.api_v1 import (
    ExportFilteredRequest,
    export_run_xlsx,
    export_run_xlsx_filtered,
)
from app.database import Base
from app.models import db_models
from app.services.schedule_export import (
    FIRST_DAY_COL,
    ROW_SUBTITLE,
    infer_run_start,
    ROW_DATE,
    ROW_DAYNUM,
    ROW_FILTER,
    ROW_FIRST_DATA,
    build_schedule_workbook,
    suggested_filename,
)

# --- sample run -------------------------------------------------------------

PLAN = SimpleNamespace(name="Nov 2025", month_days=8, current_date="2025-11-01")

COMPONENTS = [
    {"component_id": "100-A1", "name": "Lid", "quantity": 1000, "finished": 100, "mold_id": "M-1", "order_code": "A1"},
    {"component_id": "200-B1", "name": "Body", "quantity": 500, "finished": 0, "mold_id": "M-2", "order_code": "B1"},
    {"component_id": "300-C1", "name": "Hinge", "quantity": 300, "finished": 300, "mold_id": "M-3", "order_code": "C1"},
    {"component_id": "400-D1", "name": "Unscheduled", "quantity": 50, "finished": 0, "mold_id": "M-4", "order_code": "D1"},
]


def _a(day, machine, task, cid=None, qty=None, mold=None):
    return {"day": day, "machine_id": machine, "task_type": task,
            "component_id": cid, "produced_qty": qty, "mold_id": mold}


ASSIGNMENTS = [
    _a(1, "M01", "CHANGE_MOLD", "100-A1", None, "M-1"),   # 0
    _a(1, "M01", "PRODUCE", "100-A1", 300, "M-1"),        # 1
    _a(2, "M01", "PRODUCE", "100-A1", 400, "M-1"),        # 2
    _a(2, "M02", "PRODUCE", "200-B1", 200, "M-2"),        # 3
    _a(3, "M02", "WAIT"),                                 # 4
    _a(3, "M02", "PRODUCE", "200-B1", 250, "M-2"),        # 5
    _a(5, "M03", "PRODUCE", "100-A1", 200, "M-1"),        # 6
]


def _sheet(wb):
    return wb.active


def _data_rows(ws):
    rows = []
    r = ROW_FIRST_DATA
    while ws.cell(r, 1).value is not None:
        rows.append(r)
        r += 1
    return rows


def _day_headers(ws):
    cols = []
    c = FIRST_DAY_COL
    while ws.cell(ROW_DAYNUM, c).value is not None:
        cols.append(ws.cell(ROW_DAYNUM, c).value)
        c += 1
    return cols


# --- builder: unfiltered stays as it was ------------------------------------

def test_unfiltered_export_has_every_component_and_every_day():
    ws = _sheet(build_schedule_workbook(PLAN, ASSIGNMENTS, COMPONENTS))
    assert len(_data_rows(ws)) == 4
    assert _day_headers(ws) == list(range(1, 9))
    assert ws.cell(ROW_FILTER, 3).value is None


def test_unfiltered_diff_is_the_plain_formula():
    ws = _sheet(build_schedule_workbook(PLAN, ASSIGNMENTS, COMPONENTS))
    assert ws.cell(ROW_FIRST_DATA, 7).value == f"=E{ROW_FIRST_DATA}-F{ROW_FIRST_DATA}-SUM(J{ROW_FIRST_DATA}:Q{ROW_FIRST_DATA})"


# --- builder: filtered ------------------------------------------------------

def _filtered(assignments, **view):
    return build_schedule_workbook(
        PLAN, assignments, COMPONENTS,
        all_assignments=ASSIGNMENTS, filter_description="FILTERED VIEW - x", **view,
    )


def test_machine_filter_keeps_only_components_on_that_machine_and_renumbers():
    only_m02 = [a for a in ASSIGNMENTS if a["machine_id"] == "M02"]
    ws = _sheet(_filtered(only_m02))
    rows = _data_rows(ws)
    assert [ws.cell(r, 2).value for r in rows] == ["200"]        # Code of 200-B1
    assert [ws.cell(r, 1).value for r in rows] == [1]            # renumbered, not 2
    assert ws.cell(rows[0], 9).value == "M02"                    # Machine(s) from the view only


def test_machine_filter_limits_quantities_to_that_machine():
    only_m01 = [a for a in ASSIGNMENTS if a["machine_id"] == "M01"]
    ws = _sheet(_filtered(only_m01))
    r = _data_rows(ws)[0]
    day = lambda d: ws.cell(r, FIRST_DAY_COL + d - 1).value
    assert (day(1), day(2), day(5)) == (300, 400, None)          # M03's day-5 200 is out of view


def test_day_range_limits_columns_and_labels_them_with_real_days_and_dates():
    in_range = [a for a in ASSIGNMENTS if 2 <= a["day"] <= 3]
    ws = _sheet(_filtered(in_range, day_range=(2, 3)))
    assert _day_headers(ws) == [2, 3]
    assert ws.cell(ROW_DATE, FIRST_DAY_COL).value == "02/11"     # day 2 of a plan starting 1 Nov
    assert ws.cell(ROW_DATE, FIRST_DAY_COL + 1).value == "03/11"


def test_day_range_never_leaks_quantities_from_outside_it():
    # Even if the caller passes the unfiltered list, only the shown days count.
    ws = _sheet(_filtered(ASSIGNMENTS, day_range=(2, 3)))
    rows = _data_rows(ws)
    a1 = next(r for r in rows if ws.cell(r, 2).value == "100")
    assert [ws.cell(a1, FIRST_DAY_COL + i).value for i in range(2)] == [400, None]
    # ...and out-of-range data must not leak into the other columns either:
    assert ws.cell(a1, 9).value == "M01"                 # not "M01, M03" (day 5 is out of range)
    assert ws.cell(a1, 7).value.endswith("-500")         # 900 scheduled in the run - 400 shown


def test_diff_still_means_unscheduled_in_the_run_not_missing_from_the_view():
    # 100-A1: ordered 1000, finished 100, scheduled 300+400+200 = 900 in the run.
    only_m01 = [a for a in ASSIGNMENTS if a["machine_id"] == "M01"]   # shows 700 of those 900
    ws = _sheet(_filtered(only_m01))
    r = _data_rows(ws)[0]
    formula = ws.cell(r, 7).value
    assert formula.endswith("-200")                                   # the day-5 M03 production
    # Evaluate it: 1000 - 100 - (300 + 400) - 200 = 0  -> fully scheduled, not "200 short".
    shown = sum(ws.cell(r, FIRST_DAY_COL + i).value or 0 for i in range(8))
    assert 1000 - 100 - shown - 200 == 0


def test_filtered_export_notes_the_filter_under_the_title():
    ws = _sheet(_filtered([ASSIGNMENTS[3]]))
    assert ws.cell(ROW_FILTER, 3).value == "FILTERED VIEW - x"


def test_a_day_range_with_no_days_in_the_plan_is_rejected():
    with pytest.raises(ValueError):
        build_schedule_workbook(PLAN, ASSIGNMENTS, COMPONENTS,
                                all_assignments=ASSIGNMENTS, day_range=(20, 30))


def test_filtered_filename_says_so():
    assert suggested_filename(PLAN, filtered=True).endswith("_filtered.xlsx")
    assert not suggested_filename(PLAN).endswith("_filtered.xlsx")


# --- endpoints --------------------------------------------------------------

@pytest.fixture
def db():
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    session = sessionmaker(bind=engine)()
    plan = db_models.Plan(id="p1", name="Nov 2025", current_date="2025-11-01",
                          start_time="04:00", month_days=8)
    session.add(plan)
    for c in COMPONENTS:
        session.add(db_models.Component(
            plan_id="p1", component_id=c["component_id"], order_code=c["order_code"],
            name=c["name"], quantity=c["quantity"], finished=c["finished"],
            cycle_time_sec=10.0, mold_id=c["mold_id"], color="", prerequisites=[],
        ))
    session.add(db_models.Run(id="r1", plan_id="p1", run_name="Run 1", status="completed",
                              assignments=ASSIGNMENTS, unmet={}, score=0.0))
    session.commit()
    yield session
    session.close()


def _body_bytes(response):
    async def collect():
        return b"".join([chunk async for chunk in response.body_iterator])
    return asyncio.run(collect())


def _workbook(response):
    return openpyxl.load_workbook(io.BytesIO(_body_bytes(response))).active


def _filtered_request(db, **body):
    return export_run_xlsx_filtered("p1", "r1", ExportFilteredRequest(**body), revision="0", db=db)


def test_endpoint_exports_just_the_selected_assignments(db):
    resp = _filtered_request(db, assignment_indices=[3, 4, 5], description="Machine: M02")
    ws = _workbook(resp)
    assert [ws.cell(r, 2).value for r in _data_rows(ws)] == ["200"]
    assert ws.cell(ROW_FILTER, 3).value == "FILTERED VIEW - Machine: M02"
    assert "_filtered.xlsx" in resp.headers["content-disposition"]


def test_endpoint_applies_the_day_range(db):
    ws = _workbook(_filtered_request(db, assignment_indices=[2, 3, 4, 5], day_start=2, day_end=3))
    assert _day_headers(ws) == [2, 3]


def test_endpoint_with_nothing_selected_is_a_clear_409_not_an_empty_file(db):
    with pytest.raises(HTTPException) as err:
        _filtered_request(db, assignment_indices=[])
    assert err.value.status_code == 409
    assert "nothing to export" in err.value.detail


def test_endpoint_ignores_out_of_range_indices(db):
    with pytest.raises(HTTPException) as err:
        _filtered_request(db, assignment_indices=[99, -1])
    assert err.value.status_code == 409


def test_endpoint_rejects_a_day_range_outside_the_plan(db):
    with pytest.raises(HTTPException) as err:
        _filtered_request(db, assignment_indices=[1], day_start=20, day_end=30)
    assert err.value.status_code == 422


def test_endpoint_unknown_run_is_404(db):
    with pytest.raises(HTTPException) as err:
        export_run_xlsx_filtered("p1", "nope", ExportFilteredRequest(assignment_indices=[0]),
                                 revision="0", db=db)
    assert err.value.status_code == 404


def test_user_typed_search_text_can_never_become_an_excel_formula(db):
    ws = _workbook(_filtered_request(db, assignment_indices=[1], description='=HYPERLINK("http://x")'))
    cell = ws.cell(ROW_FILTER, 3)
    assert cell.data_type != "f"
    assert cell.value.startswith("FILTERED VIEW - ")


def test_the_full_download_is_unaffected_by_the_filtered_endpoint(db):
    ws = _workbook(export_run_xlsx("p1", "r1", revision="0", db=db))
    assert len(_data_rows(ws)) == 4
    assert _day_headers(ws) == list(range(1, 9))
    assert ws.cell(ROW_FILTER, 3).value is None


# --- a run keeps the start date it was generated with ----------------------------

def _dated(start: date, tasks=ASSIGNMENTS):
    """The sample tasks, each stamped with the date its day falls on."""
    from datetime import timedelta
    return [dict(t, date=(start + timedelta(days=t["day"] - 1)).isoformat()) for t in tasks]


def test_infer_run_start_reads_day_one_from_the_tasks():
    assert infer_run_start(_dated(date(2025, 11, 1))) == date(2025, 11, 1)


def test_infer_run_start_works_even_if_day_one_has_no_task():
    only_day_three = [t for t in _dated(date(2025, 11, 1)) if t["day"] == 3]
    assert infer_run_start(only_day_three) == date(2025, 11, 1)


def test_infer_run_start_crosses_month_and_year_boundaries():
    assert infer_run_start([{"day": 5, "date": "2026-01-02"}]) == date(2025, 12, 29)


def test_infer_run_start_majority_beats_a_stale_date():
    tasks = _dated(date(2025, 11, 1)) + [{"day": 9, "date": "2025-11-04"}]   # moved by hand
    assert infer_run_start(tasks) == date(2025, 11, 1)


def test_infer_run_start_falls_back_to_start_datetime_and_ignores_junk():
    assert infer_run_start([{"day": 2, "start_datetime": "2025-11-02T05:00:00"}]) == date(2025, 11, 1)
    assert infer_run_start([{"day": 1, "date": ""}, {"day": 0, "date": "2025-11-01"}, {"day": 2, "date": "nope"}]) is None
    assert infer_run_start(None) is None


def test_export_dates_follow_the_run_not_the_plans_current_start():
    # Run #1 started 1 Nov; the plan has since been moved to 2 Nov.
    plan = SimpleNamespace(name="Nov 2025", month_days=8, current_date="2025-11-02")
    ws = _sheet(build_schedule_workbook(plan, _dated(date(2025, 11, 1)), COMPONENTS))
    assert ws.cell(ROW_DATE, FIRST_DAY_COL).value == "01/11"
    assert ws.cell(ROW_DATE, FIRST_DAY_COL + 1).value == "02/11"


def test_export_title_month_follows_the_run_too():
    plan = SimpleNamespace(name="p", month_days=8, current_date="2025-11-02")
    ws = _sheet(build_schedule_workbook(plan, _dated(date(2025, 10, 30)), COMPONENTS))
    assert ws.cell(ROW_SUBTITLE, 3).value.endswith("OCTOBER 2025")


def test_filtered_export_dates_follow_the_run_as_well():
    plan = SimpleNamespace(name="p", month_days=8, current_date="2025-11-02")
    run = _dated(date(2025, 11, 1))
    ws = _sheet(build_schedule_workbook(plan, [run[3]], COMPONENTS, all_assignments=run,
                                        day_range=(2, 2), filter_description="x"))
    assert ws.cell(ROW_DATE, FIRST_DAY_COL).value == "02/11"          # day 2 of a run that began 1 Nov


def test_export_without_task_dates_falls_back_to_the_plan_start():
    plan = SimpleNamespace(name="p", month_days=8, current_date="2025-11-02")
    ws = _sheet(build_schedule_workbook(plan, ASSIGNMENTS, COMPONENTS))   # these tasks carry no dates
    assert ws.cell(ROW_DATE, FIRST_DAY_COL).value == "02/11"


def test_export_keeps_every_day_the_run_uses_even_if_the_plan_was_shortened():
    plan = SimpleNamespace(name="p", month_days=3, current_date="2025-11-01")   # run reaches day 5
    ws = _sheet(build_schedule_workbook(plan, ASSIGNMENTS, COMPONENTS))
    assert _day_headers(ws) == [1, 2, 3, 4, 5]
    a1 = next(r for r in _data_rows(ws) if ws.cell(r, 2).value == "100")
    assert ws.cell(a1, FIRST_DAY_COL + 4).value == 200                       # day-5 production is not cut off
